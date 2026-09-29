#!/usr/bin/env python3
"""
Convert a .pptx file into a folder containing a Markdown file and images.

Usage:
    $ python3 pptx2md.py lecture.pptx
    $ python3 pptx2md.py lecture.pptx -o some/dir --no-notes

Output (for lecture.pptx):
    lecture/
        lecture.md
        images/
            slide03_01.jpg
            ...

Requires: python-pptx, Pillow
"""

import argparse
import hashlib
import io
import os
import re
import sys

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.util import Emu

from PIL import Image

NS_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"

TITLE_TYPES = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE,
               PP_PLACEHOLDER.VERTICAL_TITLE}
SUBTITLE_TYPES = {PP_PLACEHOLDER.SUBTITLE}
SKIP_TYPES = {PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.FOOTER,
              PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.HEADER}

MONO_FONTS = ("courier", "consolas", "menlo", "monaco", "source code",
              "fira code", "fira mono", "lucida console", "roboto mono",
              "jetbrains mono", "sf mono", "inconsolata")

# typographic characters mapped to plain ASCII
ASCII_MAP = {
    "‘": "'", "’": "'", "“": '"', "”": '"',
    "–": "-", "—": "--", "…": "...", " ": " ",
    "→": "->", "←": "<-", "⇒": "=>", "≤": "<=",
    "≥": ">=", "≠": "!=", "×": "x", "•": "-",
    "​": "", " ": " ", "−": "-", "⋅": "*",
    "·": "*", "≈": "~=", "∞": "inf", "∑": "sum",
    "π": "pi", "Θ": "Theta", "θ": "theta", "Ω": "Omega",
    "ω": "omega", "ε": "epsilon", "λ": "lambda",
    "α": "alpha", "β": "beta", "γ": "gamma", "δ": "delta",
    "μ": "mu", "σ": "sigma", "Σ": "Sigma", "⇔": "<=>",
    "⟺": "<=>", "∀": "for all", "∃": "exists",
    "∈": "in",
}


def to_ascii(text):
    for k, v in ASCII_MAP.items():
        text = text.replace(k, v)
    return text


# ---------------------------------------------------------------- images

class ImageSaver:
    def __init__(self, img_dir, rel_dir, max_side, quality):
        self.img_dir = img_dir
        self.rel_dir = rel_dir
        self.max_side = max_side
        self.quality = quality
        self.seen = {}  # sha1 of blob -> relative path (dedupe)

    def save(self, blob, orig_ext, slide_no, index):
        digest = hashlib.sha1(blob).hexdigest()
        if digest in self.seen:
            return self.seen[digest]
        os.makedirs(self.img_dir, exist_ok=True)
        base = "slide%02d_%02d" % (slide_no, index)
        try:
            img = Image.open(io.BytesIO(blob))
            img.load()
            if img.mode in ("RGBA", "LA", "P"):
                img = img.convert("RGBA")
                bg = Image.new("RGB", img.size, (255, 255, 255))
                bg.paste(img, mask=img.split()[-1])
                img = bg
            elif img.mode != "RGB":
                img = img.convert("RGB")
            img.thumbnail((self.max_side, self.max_side))
            name = base + ".jpg"
            img.save(os.path.join(self.img_dir, name), "JPEG",
                     quality=self.quality, optimize=True, progressive=True)
        except Exception:
            # formats Pillow cannot read (e.g. EMF/WMF): keep the original
            name = base + "." + (orig_ext or "bin")
            with open(os.path.join(self.img_dir, name), "wb") as f:
                f.write(blob)
        rel = self.rel_dir + "/" + name
        self.seen[digest] = rel
        return rel


# ---------------------------------------------------------------- text

def is_mono(typeface):
    t = (typeface or "").lower()
    return any(m in t for m in MONO_FONTS)


def run_font(r_el, par_el):
    """Return (typeface, bold, italic) for an a:r, falling back to the
    paragraph default run properties."""
    face, bold, ital = None, False, False
    pPr = par_el.find(NS_A + "pPr")
    defRPr = pPr.find(NS_A + "defRPr") if pPr is not None else None
    for rpr in (defRPr, r_el.find(NS_A + "rPr")):
        if rpr is None:
            continue
        latin = rpr.find(NS_A + "latin")
        if latin is not None:
            face = latin.get("typeface")
        if rpr.get("b") is not None:
            bold = rpr.get("b") in ("1", "true")
        if rpr.get("i") is not None:
            ital = rpr.get("i") in ("1", "true")
    return face, bold, ital


def fmt_run(text, face, bold, ital):
    if not text.strip():
        return text
    lead = text[:len(text) - len(text.lstrip())]
    trail = text[len(text.rstrip()):]
    core = text.strip()
    if is_mono(face):
        if "`" not in core:
            core = "`" + core + "`"
    elif bold and ital:
        core = "***" + core + "***"
    elif bold:
        core = "**" + core + "**"
    elif ital:
        core = "*" + core + "*"
    return lead + core + trail


def script_mark(r_el, text):
    """Mark superscript runs with ^ and subscript runs with _."""
    rpr = r_el.find(NS_A + "rPr")
    try:
        base = int(rpr.get("baseline", "0")) if rpr is not None else 0
    except ValueError:
        base = 0
    core = text.strip()
    if base == 0 or not core:
        return text
    mark = "^" if base > 0 else "_"
    return mark + (core if len(core) == 1 else "(" + core + ")")


def clean(text, ascii_only):
    text = text.replace("\x0b", " ").replace(" ", " ")
    return to_ascii(text) if ascii_only else text


def paragraph_text(par, ascii_only, raw=False):
    """Text of a paragraph, walking the XML so equations are included.
    raw=True keeps whitespace and skips markdown formatting (for code)."""
    p = par._p if hasattr(par, "_p") else par
    out = []
    for child in p:
        tag = child.tag
        if tag in (NS_A + "r", NS_A + "fld"):
            t = child.find(NS_A + "t")
            text = clean(t.text or "", ascii_only) if t is not None else ""
            if raw:
                out.append(text)
            else:
                text = fmt_run(text, *run_font(child, p))
                out.append(script_mark(child, text))
        elif tag == NS_A + "br":
            out.append("\n" if raw else " ")
        else:
            math = find_math(child)
            if math is not None:
                out.append(math_md(math))
    text = "".join(out)
    if raw:
        return text.replace("\x00", "").rstrip()
    # merge adjacent markers produced by consecutive formatted runs
    text = text.replace("****", "").replace("``", "")
    text = " ".join(text.split())
    # display equations ($$...$$) go on their own lines
    return re.sub(r" ?\x00 ?", "\n\n", text).strip()


def frame_is_code(tf):
    """True when every non-empty run in the frame uses a monospace font
    and the frame has no equations."""
    seen = False
    for par in tf.paragraphs:
        if any(find_math(c) is not None for c in par._p):
            return False
        for r in par._p.findall(NS_A + "r"):
            t = r.find(NS_A + "t")
            if t is None or not (t.text or "").strip():
                continue
            if not is_mono(run_font(r, par._p)[0]):
                return False
            seen = True
    return seen


# ---------------------------------------------------------------- math

NS_M = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"
NS_A14 = "{http://schemas.microsoft.com/office/drawing/2010/main}"
NS_MC = "{http://schemas.openxmlformats.org/markup-compatibility/2006}"

MATH_SYMBOLS = {
    "Θ": "\\Theta ", "Ω": "\\Omega ", "ω": "\\omega ",
    "θ": "\\theta ", "ε": "\\varepsilon ", "λ": "\\lambda ",
    "α": "\\alpha ", "β": "\\beta ", "γ": "\\gamma ",
    "δ": "\\delta ", "π": "\\pi ", "μ": "\\mu ",
    "σ": "\\sigma ", "Σ": "\\Sigma ", "∑": "\\sum ",
    "∏": "\\prod ", "∫": "\\int ", "∞": "\\infty ",
    "≈": "\\approx ", "≤": "\\le ", "≥": "\\ge ",
    "≠": "\\ne ", "⋅": "\\cdot ", "·": "\\cdot ",
    "×": "\\times ", "→": "\\to ", "⇒": "\\Rightarrow ",
    "⟺": "\\iff ", "⇔": "\\iff ", "∀": "\\forall ",
    "∃": "\\exists ", "∈": "\\in ", "∉": "\\notin ",
    "…": "\\ldots ", "⋯": "\\cdots ", "⌊": "\\lfloor ",
    "⌋": "\\rfloor ", "⌈": "\\lceil ", "⌉": "\\rceil ",
    "−": "-", "′": "'", " ": " ", "⁡": "",
    "⁢": "", "​": "",
}

FUNC_NAMES = ("log", "lg", "ln", "lim", "max", "min", "sin", "cos", "exp")


def find_math(el):
    """Return the m:oMath / m:oMathPara element inside a14:m or an
    mc:AlternateContent wrapper, or None."""
    if el.tag == NS_A14 + "m":
        for c in el:
            if c.tag in (NS_M + "oMath", NS_M + "oMathPara"):
                return c
    if el.tag == NS_MC + "AlternateContent":
        for c in el.iter(NS_A14 + "m"):
            return find_math(c)
    return None


def m_text(s):
    return "".join(MATH_SYMBOLS.get(ch, ch) for ch in s)


def m_child(el, name):
    c = el.find(NS_M + name)
    return omml(c) if c is not None else ""


def m_val(el, prop, name, default):
    pr = el.find(NS_M + prop)
    if pr is None:
        return default
    c = pr.find(NS_M + name)
    if c is None:
        return default
    return c.get(NS_M + "val", default)


def omml(el):
    """Convert an OMML element to a LaTeX string (common constructs)."""
    tag = el.tag[len(NS_M):] if el.tag.startswith(NS_M) else None
    if tag is None:
        return ""
    if tag == "r":
        t = el.find(NS_M + "t")
        text = (t.text or "") if t is not None else ""
        rpr = el.find(NS_M + "rPr")
        if rpr is not None and rpr.find(NS_M + "nor") is not None:
            return "\\text{%s}" % text
        text = m_text(text)
        return re.sub(r"(?<!\\)(log|ln|lg|max|min|lim)", r"\\\1 ", text)
    if tag == "sSup":
        return "{%s}^{%s}" % (m_child(el, "e"), m_child(el, "sup"))
    if tag == "sSub":
        return "{%s}_{%s}" % (m_child(el, "e"), m_child(el, "sub"))
    if tag == "sSubSup":
        return "{%s}_{%s}^{%s}" % (m_child(el, "e"), m_child(el, "sub"),
                                    m_child(el, "sup"))
    if tag == "f":
        return "\\frac{%s}{%s}" % (m_child(el, "num"), m_child(el, "den"))
    if tag == "rad":
        deg = m_child(el, "deg").strip()
        if deg:
            return "\\sqrt[%s]{%s}" % (deg, m_child(el, "e"))
        return "\\sqrt{%s}" % m_child(el, "e")
    if tag == "d":
        beg = m_val(el, "dPr", "begChr", "(")
        end = m_val(el, "dPr", "endChr", ")")
        sep = m_val(el, "dPr", "sepChr", ",")
        parts = [omml(e) for e in el.findall(NS_M + "e")]
        beg = {"{": "\\{", "": "."}.get(beg, m_text(beg).strip())
        end = {"}": "\\}", "": "."}.get(end, m_text(end).strip())
        return "\\left%s %s \\right%s" % (beg, sep.join(parts), end)
    if tag in ("limLow", "limUpp"):
        base = m_child(el, "e").strip()
        lim = m_child(el, "lim")
        op = "_" if tag == "limLow" else "^"
        # bottom curly bracket drawn as a text character
        if lim.strip() == "⏟":
            return "\\underbrace{%s}" % base
        if base.endswith("⏟"):
            return "\\underbrace{%s}_{%s}" % (base[:-1], lim)
        if base in FUNC_NAMES:
            base = "\\" + base
        return "%s%s{%s}" % (base, op, lim)
    if tag == "nary":
        chr_ = m_val(el, "naryPr", "chr", "∫")
        return "%s_{%s}^{%s} %s" % (m_text(chr_).strip(), m_child(el, "sub"),
                                    m_child(el, "sup"), m_child(el, "e"))
    if tag == "func":
        name = m_child(el, "fName").strip()
        if name in FUNC_NAMES:
            name = "\\" + name
        return "%s{%s}" % (name, m_child(el, "e"))
    if tag == "groupChr":
        pos = m_val(el, "groupChrPr", "pos", "bot")
        cmd = "\\underbrace" if pos == "bot" else "\\overbrace"
        return "%s{%s}" % (cmd, m_child(el, "e"))
    if tag == "bar":
        return "\\overline{%s}" % m_child(el, "e")
    if tag == "acc":
        return "\\hat{%s}" % m_child(el, "e")
    if tag == "eqArr":
        rows = [omml(e) for e in el.findall(NS_M + "e")]
        return "\\begin{aligned}%s\\end{aligned}" % " \\\\ ".join(rows)
    if tag == "m":
        rows = []
        for mr in el.findall(NS_M + "mr"):
            rows.append(" & ".join(omml(e) for e in mr.findall(NS_M + "e")))
        return "\\begin{matrix}%s\\end{matrix}" % " \\\\ ".join(rows)
    if tag.endswith("Pr"):
        return ""
    # containers (oMath, oMathPara, e, num, den, sup, sub, lim, ...)
    return "".join(omml(c) for c in el)


def math_md(el):
    latex = " ".join(omml(el).split())
    if not latex:
        return ""
    if el.tag == NS_M + "oMathPara":
        return "\x00$$" + latex + "$$\x00"
    return "$" + latex + "$"


def bullet_kind(par, default):
    """Return 'number', 'bullet' or 'plain' for a paragraph."""
    pPr = par._p.find(NS_A + "pPr")
    if pPr is not None:
        if pPr.find(NS_A + "buNone") is not None:
            return "plain"
        if pPr.find(NS_A + "buAutoNum") is not None:
            return "number"
        if (pPr.find(NS_A + "buChar") is not None
                or pPr.find(NS_A + "buBlip") is not None):
            return "bullet"
    return default


def text_frame_md(tf, default_kind, ascii_only):
    if frame_is_code(tf):
        code = "\n".join(paragraph_text(p, ascii_only, raw=True)
                         for p in tf.paragraphs)
        return "```\n" + code.strip("\n") + "\n```"
    lines = []
    prev_kind = None
    for par in tf.paragraphs:
        text = paragraph_text(par, ascii_only)
        if not text:
            continue
        kind = bullet_kind(par, default_kind)
        indent = "    " * par.level
        if kind == "bullet":
            lines.append(indent + "- " + text)
        elif kind == "number":
            lines.append(indent + "1. " + text)
        else:
            if prev_kind in ("bullet", "number"):
                lines.append("")
            lines.append(indent + text if par.level else text)
            lines.append("")
        prev_kind = kind
    while lines and lines[-1] == "":
        lines.pop()
    return "\n".join(lines)


def table_md(table, ascii_only):
    rows = []
    for row in table.rows:
        cells = []
        for cell in row.cells:
            t = " ".join(paragraph_text(p, ascii_only)
                         for p in cell.text_frame.paragraphs).strip()
            cells.append(t.replace("|", "\\|"))
        rows.append(cells)
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    out = ["| " + " | ".join(rows[0]) + " |",
           "|" + "---|" * width]
    for r in rows[1:]:
        out.append("| " + " | ".join(r) + " |")
    return "\n".join(out)


def chart_md(chart, ascii_only):
    """Describe a chart and dump its data as a markdown table."""
    kind = str(chart.chart_type).split(" ")[0].split(".")[-1]
    title = ""
    try:
        if chart.has_title:
            title = chart.chart_title.text_frame.text.strip()
    except Exception:
        pass
    head = "**Chart** (%s)%s" % (kind.lower(), ": " + title if title else "")
    series = [s for plot in chart.plots for s in plot.series]
    if not series:
        return head
    try:
        cats = [str(c) for c in chart.plots[0].categories]
    except Exception:
        cats = []
    n = max(len(list(s.values)) for s in series)
    if len(cats) < n:
        cats = cats + [str(k + 1) for k in range(len(cats), n)]
    cols = ["category"] + [s.name or "series" for s in series]
    vals = [list(s.values) for s in series]
    rows = []
    for k in range(n):
        row = [cats[k]]
        for v in vals:
            x = v[k] if k < len(v) else None
            row.append("" if x is None else "%g" % x)
        rows.append(row)
    out = ["| " + " | ".join(cols) + " |", "|" + "---|" * len(cols)]
    out += ["| " + " | ".join(r) + " |" for r in rows]
    text = head + "\n\n" + "\n".join(out)
    return to_ascii(text) if ascii_only else text


# ---------------------------------------------------------------- shapes

def sorted_shapes(shapes):
    def key(s):
        top = s.top if s.top is not None else Emu(0)
        left = s.left if s.left is not None else Emu(0)
        return (int(top), int(left))
    return sorted(shapes, key=key)


def placeholder_type(shape):
    if not shape.is_placeholder:
        return None
    try:
        return shape.placeholder_format.type
    except Exception:
        return None


def is_title(shape):
    return placeholder_type(shape) in TITLE_TYPES


def squash(s):
    return re.sub(r"[^a-z0-9]", "", (s or "").lower())


def is_text_backdrop(pic, siblings):
    """True for a picture that is only the background of a text box: an
    exported shadow/fill image grouped with the text box, named after its
    text (Keynote exports text boxes with shadows this way)."""
    label = squash(pic.name) + " " + squash(
        pic._element.xpath("string(./*[1]/*[1]/@descr)"))
    for sib in siblings:
        if sib is pic or not getattr(sib, "has_text_frame", False):
            continue
        text = squash(sib.text_frame.text)[:20]
        if text and text in label:
            return True
    return False


def shapes_md(shapes, ctx, slide_no, blocks, in_group=False):
    shapes = list(shapes)
    for shape in sorted_shapes(shapes):
        ptype = placeholder_type(shape)
        if ptype in SKIP_TYPES or ptype in TITLE_TYPES:
            continue

        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            shapes_md(shape.shapes, ctx, slide_no, blocks, in_group=True)
            continue

        # pictures, including picture placeholders
        image = None
        try:
            image = shape.image
        except Exception:
            image = None
        if image is not None and in_group and is_text_backdrop(shape, shapes):
            continue
        if image is not None:
            ctx["img_index"] += 1
            rel = ctx["saver"].save(image.blob, image.ext, slide_no,
                                    ctx["img_index"])
            blocks.append("![%s](%s)" % (os.path.basename(rel), rel))
            continue

        if getattr(shape, "has_table", False) and shape.has_table:
            md = table_md(shape.table, ctx["ascii"])
            if md:
                blocks.append(md)
            continue

        if getattr(shape, "has_chart", False) and shape.has_chart:
            blocks.append(chart_md(shape.chart, ctx["ascii"]))
            continue

        if shape.has_text_frame:
            if ptype in SUBTITLE_TYPES:
                t = " ".join(paragraph_text(p, ctx["ascii"])
                             for p in shape.text_frame.paragraphs).strip()
                if t:
                    blocks.append("*" + t + "*")
                continue
            # body placeholders default to bullets, free text boxes to plain
            default = "bullet" if ptype is not None else "plain"
            md = text_frame_md(shape.text_frame, default, ctx["ascii"])
            if md:
                blocks.append(md)


def slide_title(slide, ascii_only):
    for shape in slide.shapes:
        if is_title(shape) and shape.has_text_frame:
            t = " ".join(paragraph_text(p, ascii_only)
                         for p in shape.text_frame.paragraphs).strip()
            if t:
                return t.replace("**", "")
    return None


# ---------------------------------------------------------------- main

def convert(pptx_path, out_root, max_side, quality, notes, ascii_only):
    name = os.path.splitext(os.path.basename(pptx_path))[0]
    out_dir = os.path.join(out_root, name)
    os.makedirs(out_dir, exist_ok=True)

    prs = Presentation(pptx_path)
    saver = ImageSaver(os.path.join(out_dir, "images"), "images",
                       max_side, quality)

    parts = []
    for i, slide in enumerate(prs.slides, start=1):
        ctx = {"saver": saver, "img_index": 0, "ascii": ascii_only}
        title = slide_title(slide, ascii_only)
        hidden = slide._element.get("show") in ("0", "false")
        tag = " (HIDDEN)" if hidden else ""
        if i == 1 and title:
            parts.append("# " + title + tag)
        else:
            parts.append("## " + (title or "Slide %d" % i) + tag)
        blocks = []
        shapes_md(slide.shapes, ctx, i, blocks)
        parts.extend(blocks)
        if notes and slide.has_notes_slide:
            nt = slide.notes_slide.notes_text_frame
            text = nt.text.strip() if nt is not None else ""
            if text:
                if ascii_only:
                    text = to_ascii(text)
                parts.append("\n".join("> " + ln if ln.strip() else ">"
                                       for ln in text.splitlines()))

    md_path = os.path.join(out_dir, name + ".md")
    with open(md_path, "w", encoding="utf-8") as f:
        f.write("\n\n".join(parts).rstrip() + "\n")
    return md_path, len(saver.seen)


def main():
    ap = argparse.ArgumentParser(
        description="Convert a .pptx into a folder with Markdown + images.")
    ap.add_argument("pptx", help="input .pptx file")
    ap.add_argument("-o", "--outdir", default=None,
                    help="where to create the output folder "
                         "(default: next to the .pptx)")
    ap.add_argument("--max-side", type=int, default=800,
                    help="max image width/height in pixels (default 800)")
    ap.add_argument("--quality", type=int, default=60,
                    help="JPEG quality 1-95 (default 60)")
    ap.add_argument("--no-notes", dest="notes", action="store_false",
                    help="omit speaker notes (included as blockquotes "
                         "by default)")
    ap.add_argument("--keep-unicode", action="store_true",
                    help="do not convert smart quotes, dashes, arrows "
                         "to ASCII")
    args = ap.parse_args()

    if not os.path.isfile(args.pptx):
        sys.exit("error: file not found: " + args.pptx)
    out_root = args.outdir or os.path.dirname(os.path.abspath(args.pptx))
    md_path, n_img = convert(args.pptx, out_root, args.max_side,
                             args.quality, args.notes,
                             not args.keep_unicode)
    print("wrote %s (%d images)" % (md_path, n_img))


if __name__ == "__main__":
    main()
